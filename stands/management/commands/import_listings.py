"""Import stands from listings.xlsx into the database.

Usage:
    python manage.py import_listings [path/to/listings.xlsx] [--publish] [--overwrite]

- Rows whose Notes contain 'EXAMPLE ROW' are skipped.
- Imports as DRAFT unless --publish is given, so everything gets a review
  pass in the admin before going live.
- NON-DESTRUCTIVE by default: matched by slugified name; stands that already
  exist in the DB are left completely untouched (admin edits, status, hours,
  payment methods, categories all preserved). Only NEW rows are created.
- --overwrite restores the old behavior (sheet wins, hours rebuilt) for ALL
  rows. Use only when the sheet really is the source of truth, e.g. first
  provisioning. Reports/claims/photos are never touched either way.
- Hours like '6:30am-sold out' get an ESTIMATED close time (open + 8h,
  capped 8pm) so the open-now filter works; flagged in irregular_hours_note.
"""
import datetime
import re

from django.core.management.base import BaseCommand, CommandError
from django.utils.text import slugify

from stands.models import Category, PaymentMethod, Stand, WeeklyHours

DAY_COLS = ['Hours Mon', 'Hours Tue', 'Hours Wed', 'Hours Thu',
            'Hours Fri', 'Hours Sat', 'Hours Sun']

TYPE_MAP = {
    'bake stand': Stand.LocationType.BAKE_STAND,
    'food truck': Stand.LocationType.FOOD_TRUCK,
    'pop-up': Stand.LocationType.POPUP,
    'farm stand': Stand.LocationType.FARM_STAND,
}

TIME_RE = re.compile(r'(\d{1,2})(?::(\d{2}))?\s*(am|pm)?', re.I)


def parse_time(s):
    m = TIME_RE.match(s.strip())
    if not m:
        return None
    h, mins, ap = int(m.group(1)), int(m.group(2) or 0), (m.group(3) or '').lower()
    if ap == 'pm' and h != 12:
        h += 12
    if ap == 'am' and h == 12:
        h = 0
    if h > 23 or mins > 59:
        return None
    return datetime.time(h, mins)


def parse_range(s):
    """'7am-2pm' -> (time, time, False); '6:30am-sold out' -> (time, est, True)."""
    parts = re.split(r'\s*[-–]\s*', s.strip(), maxsplit=1)
    if not parts or not parts[0]:
        return None
    open_t = parse_time(parts[0])
    if open_t is None:
        return None
    close_t = parse_time(parts[1]) if len(parts) > 1 else None
    estimated = False
    if close_t is None:  # 'sold out', 'SO', missing, etc.
        est_h = min(open_t.hour + 8, 20)
        close_t = datetime.time(est_h, open_t.minute)
        estimated = True
    return open_t, close_t, estimated


def clean_handle(s):
    return (s or '').strip().lstrip('@')


class Command(BaseCommand):
    help = 'Import stands from the listings.xlsx data-collection sheet.'

    def add_arguments(self, parser):
        parser.add_argument('path', nargs='?', default='listings.xlsx')
        parser.add_argument('--publish', action='store_true',
                            help='Import as published instead of draft.')
        parser.add_argument('--overwrite', action='store_true',
                            help='DANGER: update existing stands from the sheet, '
                                 'overwriting admin edits and rebuilding hours.')

    def handle(self, *args, **opts):
        try:
            from openpyxl import load_workbook
        except ImportError:
            raise CommandError('openpyxl is required: pip install openpyxl')

        try:
            wb = load_workbook(opts['path'], data_only=True)
        except FileNotFoundError:
            raise CommandError(f"File not found: {opts['path']}")
        ws = wb['Listings'] if 'Listings' in wb.sheetnames else wb.active

        hdr = {str(c.value).strip(): i for i, c in enumerate(ws[1]) if c.value}
        get = lambda row, col: (str(row[hdr[col]]).strip()
                                if col in hdr and row[hdr[col]] is not None else '')

        status = Stand.Status.PUBLISHED if opts['publish'] else Stand.Status.DRAFT
        created = updated = skipped = untouched = 0

        for row in ws.iter_rows(min_row=2, values_only=True):
            name = get(row, 'Name')
            if not name:
                continue
            if 'EXAMPLE ROW' in get(row, 'Notes').upper():
                skipped += 1
                continue

            loc_type = TYPE_MAP.get(get(row, 'Location Type').lower(),
                                    Stand.LocationType.BAKE_STAND)
            attendance = (Stand.Attendance.UNATTENDED
                          if get(row, 'Attendance').lower().startswith('unatt')
                          else Stand.Attendance.ATTENDED)

            lat, lng = None, None
            try:
                lat = round(float(get(row, 'Latitude')), 6)
                lng = round(float(get(row, 'Longitude')), 6)
            except (ValueError, TypeError):
                pass

            irregular = get(row, 'Irregular Hours Note')

            if not opts['overwrite'] and Stand.objects.filter(
                    slug=slugify(name)).exists():
                untouched += 1
                continue

            stand, was_created = Stand.objects.update_or_create(
                slug=slugify(name),
                defaults=dict(
                    name=name,
                    location_type=loc_type,
                    description=get(row, 'Description / What They Sell'),
                    street_address=get(row, 'Street Address'),
                    latitude=lat, longitude=lng,
                    # Sheet coords are human-researched — protect from geocoder.
                    coords_source=(Stand.CoordsSource.ADMIN
                                   if lat is not None else ''),
                    attendance=attendance,
                    phone=get(row, 'Phone'),
                    instagram=clean_handle(get(row, 'Instagram')),
                    facebook=clean_handle(get(row, 'Facebook')),
                    tiktok=clean_handle(get(row, 'TikTok')),
                    website=get(row, 'Website'),
                    email=get(row, 'Email'),
                    internal_notes=' | '.join(filter(None, [
                        get(row, 'Notes'), get(row, 'Photo Sources / Links')])),
                    status=status,
                ))

            # hours: rebuild from sheet
            stand.weekly_hours.all().delete()
            estimated_days = []
            for weekday, col in enumerate(DAY_COLS):
                raw = get(row, col)
                if not raw:
                    continue
                parsed = parse_range(raw)
                if parsed is None:
                    irregular = ' '.join(filter(None, [irregular, f'{col[6:]}: {raw}.']))
                    continue
                open_t, close_t, est = parsed
                WeeklyHours.objects.create(
                    stand=stand, weekday=weekday, open_time=open_t, close_time=close_t)
                if est:
                    estimated_days.append(col[6:])
            if estimated_days:
                note = f"Sells out (close times estimated: {', '.join(estimated_days)})"
                irregular = ' '.join(filter(None, [irregular, note]))
            stand.irregular_hours_note = irregular[:300]
            stand.save(update_fields=['irregular_hours_note'])

            # payment methods
            pm_raw = get(row, 'Payment Options')
            if pm_raw:
                pms = []
                for p in re.split(r'[,/]', pm_raw):
                    p = p.strip()
                    if p:
                        pm, _ = PaymentMethod.objects.get_or_create(name=p.title()
                            if p.lower() not in ('paypal', 'cashapp') else
                            {'paypal': 'PayPal', 'cashapp': 'CashApp'}[p.lower()])
                        pms.append(pm)
                stand.payment_methods.set(pms)

            # categories
            cat_raw = get(row, 'Categories / Tags')
            if cat_raw:
                cats = []
                for cname in cat_raw.split(','):
                    cname = cname.strip()
                    if cname:
                        cat, _ = Category.objects.get_or_create(
                            slug=slugify(cname), defaults={'name': cname.capitalize()})
                        cats.append(cat)
                stand.categories.set(cats)

            created += was_created
            updated += not was_created

        self.stdout.write(self.style.SUCCESS(
            f'Imported: {created} created, {updated} updated (--overwrite), '
            f'{untouched} existing left untouched, {skipped} skipped (example rows). '
            f'Status: {status}.'))
        if not opts['publish']:
            self.stdout.write('Review drafts in the admin, then publish — or rerun with --publish.')
